from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

ROOT=Path('/Users/macbookpro/Documents/Playground/cancha-central')
OUT=ROOT/'Entrega_2_Cancha_Central.docx'
MOCKUPS=Path('/Users/macbookpro/Documents/Playground/entrega1_assets')

doc=Document(); sec=doc.sections[0]
for attr in ('top_margin','bottom_margin','left_margin','right_margin'): setattr(sec,attr,Inches(1))
styles=doc.styles
styles['Normal'].font.name='Times New Roman';styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman');styles['Normal'].font.size=Pt(12);styles['Normal'].paragraph_format.line_spacing=2
for name,size in [('Title',16),('Heading 1',14),('Heading 2',12)]:
    s=styles[name];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.bold=True;s.font.color.rgb=RGBColor(0,0,0);s.paragraph_format.line_spacing=2

def p(text='',style=None,align=None,indent=True,before=0,after=0):
    q=doc.add_paragraph(style=style);q.paragraph_format.line_spacing=2;q.paragraph_format.space_before=Pt(before);q.paragraph_format.space_after=Pt(after)
    if indent and style is None and align is None and text:q.paragraph_format.first_line_indent=Inches(.5)
    if align:q.alignment=align
    q.add_run(text);return q
def page():doc.add_page_break()
def fig(path,caption):
    q=doc.add_picture(str(path),width=Inches(6.45));doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    p(caption,align=WD_ALIGN_PARAGRAPH.CENTER,indent=False,before=4,after=8)
def code(text):
    q=doc.add_paragraph();q.paragraph_format.left_indent=Inches(.3);q.paragraph_format.right_indent=Inches(.25);q.paragraph_format.space_after=Pt(7);q.paragraph_format.line_spacing=1
    run=q.add_run(text);run.font.name='Courier New';run.font.size=Pt(8);run.font.color.rgb=RGBColor(35,45,55)
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F2F4F5');q._p.get_or_add_pPr().append(shade)
def table(rows):
    t=doc.add_table(rows=1,cols=2);t.style='Table Grid';hdr=t.rows[0].cells;hdr[0].text='Elemento';hdr[1].text='Descripción'
    for a,b in rows:c=t.add_row().cells;c[0].text=a;c[1].text=b
    for ri,row in enumerate(t.rows):
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in cell.paragraphs:
                para.paragraph_format.line_spacing=1.1
                for run in para.runs:run.font.name='Times New Roman';run.font.size=Pt(10);run.font.bold=(ri==0)
    return t

# Cover
p('Entrega 2 Prototipo Funcional de Aplicación Web',style='Title',align=WD_ALIGN_PARAGRAPH.CENTER,before=150)
p('Cancha Central Plataforma Web de Noticias Deportivas',align=WD_ALIGN_PARAGRAPH.CENTER,indent=False,before=38)
p('David Santiago Rodríguez Cruz',align=WD_ALIGN_PARAGRAPH.CENTER,indent=False,before=28)
p('Politécnico Grancolombiano',align=WD_ALIGN_PARAGRAPH.CENTER,indent=False)
p('John Olarte Ramos',align=WD_ALIGN_PARAGRAPH.CENTER,indent=False)
p('26 de septiembre de 2026',align=WD_ALIGN_PARAGRAPH.CENTER,indent=False)
page()

p('Tabla de contenido',style='Heading 1',align=WD_ALIGN_PARAGRAPH.CENTER)
for line in ['Introducción .......................................................... 3','Objetivos .................................................................... 3','Implementación del prototipo ....................................... 4','Estructura del código .................................................. 4','Maquetación de la Entrega 1 ......................................... 5','Funcionalidades implementadas ..................................... 9','Fragmentos de código fuente ........................................ 10','Repositorio, conclusiones y referencias ......................... 11']:
    p(line,indent=False,before=6)
page()

p('Introducción',style='Heading 1')
p('Este documento presenta la segunda entrega de Cancha Central, una plataforma web de noticias deportivas. El propósito de esta fase fue transformar la propuesta visual de la Entrega 1 en un prototipo funcional que permitiera explorar noticias, consultar su detalle, guardar favoritos, enviar mensajes y administrar publicaciones de forma local.')
p('La versión desarrollada utiliza HTML para la estructura semántica, CSS para el diseño responsive y JavaScript para el renderizado dinámico y las interacciones. Las noticias se obtienen desde un archivo JSON local y las acciones del usuario se conservan en localStorage. MDN Web Docs (2025a) señala que HTML organiza el contenido de una página, mientras que JavaScript permite dotar de comportamiento a la interfaz (MDN Web Docs, 2025b).')
p('Objetivos',style='Heading 1',before=18)
p('Objetivo general',style='Heading 2',before=10)
p('Implementar una primera versión funcional de Cancha Central que corresponda con los mockups previamente definidos y evidencie el uso integrado de HTML, CSS y JavaScript.')
p('Objetivos específicos',style='Heading 2',before=10)
for x in ['Construir las páginas Home, Listado, Detalle, Favoritos, Contacto y Administración.','Renderizar noticias mediante un archivo JSON local.','Implementar favoritos persistentes mediante localStorage.','Validar el formulario de contacto y mostrar un mensaje de confirmación.','Implementar un mini CRUD para crear y eliminar noticias locales.']:p(x,indent=False,before=4)
page()

p('Implementación del prototipo',style='Heading 1')
p('El prototipo mantiene la identidad visual definida en la primera entrega: azul marino para la navegación y la información editorial, verde para acciones principales, amarillo como color de énfasis y marfil como fondo. El diseño responde a pantallas de computador y dispositivos móviles mediante media queries en CSS.')
p('Tecnologías utilizadas',style='Heading 2',before=12)
table([('HTML5','Estructura semántica de las seis páginas y formularios accesibles.'),('CSS3','Sistema visual, grid, tarjetas, diseño responsive y estados de interacción.'),('JavaScript','Carga dinámica del JSON, filtros, favoritos, validaciones y CRUD.'),('JSON local','Fuente inicial de noticias ubicada en data/noticias.json.'),('localStorage','Persistencia de favoritos y noticias creadas desde el panel editorial.')])
p('Estructura del código',style='Heading 1',before=18)
p('La organización separa responsabilidades para facilitar la lectura y el mantenimiento del código. Los archivos de presentación se ubican en la raíz, los estilos en css, la lógica compartida en js, los datos en data y los recursos gráficos en assets.')
code('cancha-central/\n├── index.html | noticias.html | detalle.html | favoritos.html\n├── contacto.html | administrar.html\n├── css/styles.css\n├── js/app.js\n├── data/noticias.json\n└── assets/noticia-*.svg')
page()

p('Maquetación de la Entrega 1',style='Heading 1')
p('Las siguientes pantallas corresponden a la propuesta visual utilizada como guía durante el desarrollo. La implementación funcional conserva la navegación, las tarjetas de noticias, los llamados a la acción y la identidad visual planteada en Figma.')
p('Diseño de Figma: https://www.figma.com/design/t0vgghM0RM6XpV8sWRljBA/Cancha-Central-%E2%80%94-Mockups?node-id=0-1',indent=False,before=4)
fig(MOCKUPS/'home.png','Figura 1. Maquetación de la página Home de Cancha Central.')
fig(MOCKUPS/'list.png','Figura 2. Maquetación del listado de noticias.')
page()
p('Maquetación de la Entrega 1',style='Heading 1')
fig(MOCKUPS/'detail.png','Figura 3. Maquetación del detalle de noticia.')
page()
p('Maquetación de la Entrega 1',style='Heading 1')
fig(MOCKUPS/'contact.png','Figura 4. Maquetación de la página de contacto.')
page()

p('Funcionalidades implementadas',style='Heading 1')
table([('Noticias dinámicas','noticias.html consulta data/noticias.json y genera tarjetas con imagen, categoría, título, descripción y enlace.'),('Búsqueda y filtros','El listado filtra las noticias por texto y por categoría sin recargar la página.'),('Detalle','detalle.html recupera el identificador de la noticia desde la URL y muestra su contenido completo.'),('Favoritos','El usuario agrega o elimina favoritos. La lista se conserva mediante localStorage, tecnología descrita por MDN Web Docs (2025c).'),('Contacto','El formulario valida campos obligatorios, formato de correo y longitud mínima del mensaje antes de confirmar el envío.'),('Mini CRUD','administrar.html permite crear noticias y eliminarlas; los cambios se almacenan localmente para la sesión posterior.')])
p('La interfaz incluye un menú de seis páginas, un pie de página uniforme y adaptación responsive. La estructura se diseñó para que el código pueda evolucionar a componentes de Angular en la siguiente entrega.')
page()

p('Fragmentos de código fuente',style='Heading 1')
p('Carga de noticias desde JSON',style='Heading 2',before=10)
code("baseNews = await fetch('data/noticias.json').then(r => r.json());\nfunction newsCard(item) {\n  return `<article class=\"card\"> ... ${item.title} ... </article>`;\n}")
p('Persistencia de favoritos',style='Heading 2',before=10)
code("function toggleFavorite(id) {\n  const items = favorites();\n  setLocal(KEYS.favorites, items.includes(id)\n    ? items.filter(x => x !== id) : [...items, id]);\n}")
p('Validación del formulario',style='Heading 2',before=10)
code("form.addEventListener('submit', e => {\n  e.preventDefault();\n  const bad = !field.validity.valid;\n  field.classList.toggle('invalid', bad);\n  if (valid) success.hidden = false;\n});")
p('Estos fragmentos se encuentran comentados y centralizados en js/app.js. El repositorio local contiene el código fuente completo y el archivo README con las instrucciones para ejecutar el servidor local.')
page()

p('Repositorio del proyecto',style='Heading 1')
p('El código fuente se encuentra organizado en un repositorio Git local llamado cancha-central. La URL del repositorio público se agregará después de su publicación en la cuenta de GitHub del estudiante. El proyecto incluye un archivo README con instrucciones de ejecución mediante un servidor local.')
p('Conclusiones',style='Heading 1',before=18)
p('La Entrega 2 demuestra la transición de la maquetación hacia un prototipo web funcional. El sitio conserva la identidad de Cancha Central y resuelve los recorridos principales de un usuario: descubrir noticias, filtrarlas, consultar su detalle, guardarlas como favoritas y contactar al equipo editorial.')
p('El uso de un archivo JSON y localStorage permite evidenciar una gestión básica de datos sin requerir un servidor. Esta arquitectura facilita la evolución futura hacia Angular, componentes reutilizables y un despliegue público en la entrega final.')
p('Referencias',style='Heading 1',align=WD_ALIGN_PARAGRAPH.CENTER,before=18)
refs=['MDN Web Docs. (2025a). HTML: HyperText Markup Language. Mozilla. https://developer.mozilla.org/en-US/docs/Web/HTML','MDN Web Docs. (2025b). JavaScript. Mozilla. https://developer.mozilla.org/en-US/docs/Web/JavaScript','MDN Web Docs. (2025c). Window: localStorage property. Mozilla. https://developer.mozilla.org/en-US/docs/Web/API/Window/localStorage','Figma. (s. f.). Figma. https://www.figma.com/']
for r in refs:
    q=p(r,indent=False,before=6);q.paragraph_format.left_indent=Inches(.5);q.paragraph_format.first_line_indent=Inches(-.5)

for section in doc.sections:
    header=section.header.paragraphs[0];header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');header._p.append(fld)
doc.save(OUT);print(OUT)
